"""
Document Loader Module for SmartHire GenAI.

This module handles loading, validating, and extracting clean text from resume files
in PDF (.pdf) and Microsoft Word (.docx) formats. It forms the initial stage in the
candidate evaluation pipeline:
    File on disk -> validate -> extract text -> normalize & clean -> return structured result

Supported Formats:
    - .pdf  (extracted using pypdf)
    - .docx (extracted using python-docx)

Extraction Approach & Features:
    1. Strict Pre-Extraction Validation:
       - Checks file existence, non-empty paths, and supported file extensions.
       - Rejects 0-byte corrupted files.
       - Enforces configured maximum file size limits (MAX_RESUME_SIZE_BYTES from src.config).
    2. Format-Specific Stream Extraction:
       - PDF: Uses pypdf.PdfReader with space_width=300.0 to prevent artificial word breaks
         caused by wide-kerning font families (e.g. "Technology" instead of "T echnology").
       - DOCX: Iterates paragraph XML structures in natural document order.
    3. Encryption & Security Guard:
       - Explicitly detects password-protected / encrypted PDFs and returns actionable errors.
    4. Text Cleaning & Whitespace Normalization (clean_extracted_text):
       - Strips non-standard decorative font icon artifacts (e.g., telephone/envelope glyphs)
         while strictly preserving phone numbers, emails, URLs, dates, and technical punctuation.
       - Safely unwraps soft-wrapped bullet points into coherent single lines without merging
         headings, dates, or adjacent project entries.
       - Normalizes multi-space runs and excessive blank lines.
    5. Incomplete / Scanned Text Flagging:
       - Configurable minimum extracted-text threshold (default: 50 characters) flags suspiciously
         short extractions (is_suspiciously_short) without crashing on valid short snippets.
    6. Overlapping Text Chunking (chunk_text):
       - Sliding window chunker with configurable chunk size and overlap for embedding and LLM steps.

Architectural Limitations:
    - Scanned / Image-only Documents:
      This loader performs native text-stream extraction. It does NOT include an OCR engine.
      Scanned paper resumes or image-only PDFs with no selectable text will yield 0 or very few
      characters and will be flagged as empty or suspiciously short.
    - Complex Multi-Column Layouts:
      pypdf extracts text according to internal PDF object order. Documents with intricate
      side-by-side columnar layouts may interleave text unless designed with standard flow tags.
    - Embedded DOCX Objects & Tables:
      The default DOCX extractor captures body paragraphs. Embedded non-standard shapes or
      multi-nested tables require specialized cell processing if not represented as paragraphs.
    - Password Protection:
      Encrypted PDFs cannot be read without credentials and will raise a clear RuntimeError.

Usage:
    from src.parsing.loader import load_document, chunk_text, clean_extracted_text

    result = load_document("path/to/resume.pdf")
    print(result["text"])                    # Normalized extracted text
    print(result["file_type"])               # "pdf" or "docx"
    print(result["metadata"]["num_pages"])   # Total page count

    chunks = chunk_text(result["text"], chunk_size=1000, overlap=200)
"""

import logging
from pathlib import Path
import re

from src.config import MAX_RESUME_SIZE_BYTES, MIN_EXTRACTED_TEXT_CHARS

# ---------------------------------------------------------------------------
# Logger setup
# ---------------------------------------------------------------------------
logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Constants & Regex Patterns
# ---------------------------------------------------------------------------
SUPPORTED_EXTENSIONS = {".pdf", ".docx"}

# Matches lines that start with standard resume bullet glyphs or numbering
BULLET_START_PATTERN = re.compile(
    r"^\s*([\u2022*\u25aa\u25ab\u2023\u25cf\u25cb\u25e6\u2219\u2713\u2714\u25b8\u2043]|\-\s|\d+[\.\)])"
)

# Matches lines that appear to be standalone all-caps section headings
HEADING_PATTERN = re.compile(r"^[A-Z0-9\s&/,\-\u2013\u2014]{3,}$")

# Matches dates or date ranges (e.g., "Jan 2025 - Apr 2025", "2023 - 2027", "November 2024")
DATE_PATTERN = re.compile(
    r"(?:(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*|\d{4})"
    r"\s*[\-\u2013\u2014to]+\s*"
    r"(?:(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*|\d{4}|Present|Current)"
    r"|\b(?:19|20)\d{2}\b",
    re.IGNORECASE,
)

# Known decorative font icon artifacts from resume templates
# (\ue000-\uf8ff: Private Use Area, \ufffd: replacement char,
# \u2640/\u2642: gender symbols used for icons, \u260e/\u2709: phone/mail dingbats, \ufe0f: variation selector)
DECORATIVE_ICON_PATTERN = re.compile(r"[\ue000-\uf8ff\ufffd\u2640\u2642\u260e\u2709\ufe0f]")


# ===========================================================================
# TEXT CLEANING & NORMALIZATION
# ===========================================================================

def is_continuation_line(prev_line: str, curr_line: str) -> bool:
    """
    Determine whether curr_line is a soft-wrapped continuation of a bullet point on prev_line.

    Safety rules:
    - If curr_line begins with a bullet glyph or number, it is a new bullet item.
    - If curr_line matches an all-caps section heading, it is a section boundary.
    - If curr_line has a pipe separator (|) or date pattern, it is a job/project title line.
    - If prev_line ends with terminal punctuation ('.', ':', ';') AND curr_line begins
      with an uppercase letter, it is a distinct sentence/entry, not a wrapped line.
    - Otherwise, curr_line is a soft-wrapped continuation and can safely be joined.
    """
    if BULLET_START_PATTERN.match(curr_line):
        return False
    if HEADING_PATTERN.match(curr_line) and len(curr_line) > 2:
        return False
    if "|" in curr_line or DATE_PATTERN.search(curr_line):
        return False
    if prev_line.endswith((".", ":", ";")) and curr_line[0].isupper():
        return False
    return True


def clean_extracted_text(text: str) -> str:
    """
    Normalize extracted text and clean known extraction artifacts.

    Transformations performed:
    1. Standardizes line breaks (\\r\\n -> \\n, \\r -> \\n).
    2. Strips invisible zero-width and control characters.
    3. Strips decorative font icon artifacts while strictly preserving all letters,
       numbers, punctuation (., :, +, -, /, @, #, etc.), technical skills (C++, C#, .NET),
       URLs, and bullet points.
    4. Safely unwraps soft-wrapped bullet point lines without merging headings,
       dates, or distinct entries.
    5. Normalizes inline whitespace (tabs and multiple spaces -> single space).
    6. Collapses excessive consecutive blank lines (3+ newlines -> \\n\\n).

    Args:
        text: Raw extracted text string.

    Returns:
        Cleaned, normalized text string.
    """
    if not text:
        return ""

    # Step 1: Normalize line endings
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # Step 2: Strip invisible zero-width characters
    text = re.sub(r"[\u200b\u200c\u200d\ufeff]", "", text)

    # Step 3: Strip non-standard decorative font icon artifacts
    text = DECORATIVE_ICON_PATTERN.sub(" ", text)

    # Step 4: Clean line by line and unwrap soft-wrapped bullet points
    raw_lines = text.split("\n")
    cleaned_lines: list[str] = []
    in_bullet = False

    for line in raw_lines:
        line_stripped = re.sub(r"[ \t\u00a0]+", " ", line).strip()
        if not line_stripped:
            in_bullet = False
            cleaned_lines.append("")
            continue

        is_bullet = bool(BULLET_START_PATTERN.match(line_stripped))

        if is_bullet:
            in_bullet = True
            cleaned_lines.append(line_stripped)
        elif in_bullet and is_continuation_line(cleaned_lines[-1], line_stripped):
            # Continuation of previous bullet: unwrap onto previous line
            cleaned_lines[-1] = cleaned_lines[-1] + " " + line_stripped
        else:
            in_bullet = False
            cleaned_lines.append(line_stripped)

    # Step 4: Collapse consecutive blank lines and strip margins
    result = "\n".join(cleaned_lines)
    result = re.sub(r"\n{3,}", "\n\n", result).strip()
    return result


# ===========================================================================
# VALIDATION
# ===========================================================================

def validate_file(file_path: str) -> Path:
    """
    Validate that a file is suitable for text extraction.

    Checks performed (in order):
    1. Input is not None, empty string, or whitespace-only.
    2. File exists on disk.
    3. File extension is .pdf or .docx.
    4. File is not empty (0 bytes).
    5. File size does not exceed MAX_RESUME_SIZE_BYTES (from src.config).

    Args:
        file_path: Path to the file as a string.

    Returns:
        A pathlib.Path object pointing to the validated file.

    Raises:
        ValueError: If input is empty, file type is unsupported, file is
                     empty, or file exceeds the size limit.
        FileNotFoundError: If the file does not exist.
    """
    if not file_path or not str(file_path).strip():
        raise ValueError("File path cannot be empty.")

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    extension = path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: '{extension}'. "
            f"Supported types: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    file_size = path.stat().st_size
    if file_size == 0:
        raise ValueError(f"File is empty (0 bytes): {path}")

    if file_size > MAX_RESUME_SIZE_BYTES:
        size_mb = file_size / (1024 * 1024)
        limit_mb = MAX_RESUME_SIZE_BYTES / (1024 * 1024)
        raise ValueError(
            f"File too large: {size_mb:.1f} MB "
            f"(limit: {limit_mb:.0f} MB): {path}"
        )

    logger.debug("File validated: %s (%d bytes)", path, file_size)
    return path


# ===========================================================================
# PDF EXTRACTION
# ===========================================================================

def extract_text_from_pdf(file_path: Path) -> dict:
    """
    Extract and clean text from a PDF file using pypdf.

    How it works:
    1. Verifies encryption status using reader.is_encrypted.
    2. Iterates through Page objects in sequential document order.
    3. Extracts text per page using space_width=300.0 to prevent artificial word breaks.
    4. Cleans and normalizes the full text using clean_extracted_text().

    Args:
        file_path: Path object pointing to a .pdf file.

    Returns:
        A dict with:
        - "text": Cleaned full extracted text
        - "pages": Cleaned per-page text list
        - "num_pages": Total page count

    Raises:
        RuntimeError: If the PDF cannot be opened, is encrypted, has no pages,
                      or yields 0 characters of text.
    """
    try:
        from pypdf import PdfReader
    except ImportError:
        raise RuntimeError(
            "pypdf is required for PDF extraction. "
            "Install it with: pip install pypdf"
        )

    try:
        reader = PdfReader(str(file_path))
    except Exception as exc:
        raise RuntimeError(
            f"Failed to open PDF file: {file_path}. Error: {exc}"
        ) from exc

    # Check for password protection / encryption
    if getattr(reader, "is_encrypted", False):
        raise RuntimeError(
            f"PDF file is encrypted or password-protected: {file_path}. "
            f"Please decrypt or remove password protection before uploading."
        )

    num_pages = len(reader.pages)
    if num_pages == 0:
        raise RuntimeError(f"PDF has no pages: {file_path}")

    pages_text = []
    for i, page in enumerate(reader.pages):
        try:
            # space_width=300.0 prevents artificial spaces between letters with
            # wide font kerning/overhangs (e.g. "Technology" instead of "T echnology").
            text = page.extract_text(space_width=300.0) or ""
            if not text.strip():
                # Fallback to standard extraction if space_width yielded nothing
                text = page.extract_text() or ""
        except Exception as exc:
            logger.warning(
                "Failed to extract text from page %d of %s: %s",
                i + 1, file_path, exc,
            )
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""
        pages_text.append(clean_extracted_text(text))

    full_text = "\n\n".join(p for p in pages_text if p.strip()).strip()

    if not full_text:
        raise RuntimeError(
            f"PDF extraction produced no text. This document appears to be a scanned image or photo "
            f"without selectable text: {file_path}. Please upload a standard text-based PDF or DOCX file."
        )

    logger.debug(
        "PDF extracted: %d pages, %d characters from %s",
        num_pages, len(full_text), file_path,
    )

    return {
        "text": full_text,
        "pages": pages_text,
        "num_pages": num_pages,
    }


# ===========================================================================
# DOCX EXTRACTION
# ===========================================================================

def extract_text_from_docx(file_path: Path) -> dict:
    """
    Extract and clean text from a DOCX file using python-docx.

    How it works:
    1. Document() opens the .docx package.
    2. Iterates body elements in natural document order, extracting text from
       both paragraphs and tables (including multi-column and grid layouts).
    3. Normalizes and cleans the joined text using clean_extracted_text().

    Args:
        file_path: Path object pointing to a .docx file.

    Returns:
        A dict with:
        - "text": Cleaned full extracted text
        - "num_paragraphs": Count of non-empty paragraphs and content units

    Raises:
        RuntimeError: If the DOCX cannot be opened or yields no text.
    """
    try:
        from docx import Document
        from docx.text.paragraph import Paragraph
        from docx.table import Table
    except ImportError:
        raise RuntimeError(
            "python-docx is required for DOCX extraction. "
            "Install it with: pip install python-docx"
        )

    try:
        doc = Document(str(file_path))
    except Exception as exc:
        raise RuntimeError(
            f"Failed to open DOCX file: {file_path}. Error: {exc}"
        ) from exc

    # Natural document order extraction (paragraphs + tables)
    content_blocks = []
    paragraph_count = 0

    try:
        for element in doc.element.body:
            if element.tag.endswith("p"):
                p = Paragraph(element, doc)
                text = p.text.strip()
                if text:
                    paragraph_count += 1
                    content_blocks.append(text)
            elif element.tag.endswith("tbl"):
                tbl = Table(element, doc)
                for row in tbl.rows:
                    row_texts = []
                    seen_cells = set()
                    for cell in row.cells:
                        if cell._tc not in seen_cells:
                            seen_cells.add(cell._tc)
                            cell_text = cell.text.strip()
                            if cell_text:
                                row_texts.append(cell_text)
                    if row_texts:
                        content_blocks.append(" | ".join(row_texts))
                        paragraph_count += len(row_texts)
    except Exception as order_exc:
        logger.warning(
            "Natural order extraction had an issue on %s (%s); falling back to paragraphs + tables.",
            file_path, order_exc,
        )

    # Fallback to standard doc.paragraphs if body iteration produced nothing
    if not content_blocks:
        paragraphs = [para.text for para in doc.paragraphs]
        content_blocks = [p.strip() for p in paragraphs if p.strip()]
        paragraph_count = len(content_blocks)
        # Also append table content in fallback
        for tbl in doc.tables:
            for row in tbl.rows:
                seen_cells = set()
                row_texts = []
                for cell in row.cells:
                    if cell._tc not in seen_cells:
                        seen_cells.add(cell._tc)
                        cell_text = cell.text.strip()
                        if cell_text:
                            row_texts.append(cell_text)
                if row_texts:
                    content_blocks.append(" | ".join(row_texts))
                    paragraph_count += len(row_texts)

    raw_text = "\n".join(content_blocks).strip()
    full_text = clean_extracted_text(raw_text)

    if not full_text:
        raise RuntimeError(
            f"DOCX extraction produced no text: {file_path}"
        )

    logger.debug(
        "DOCX extracted: %d non-empty paragraphs/cells, %d characters from %s",
        paragraph_count, len(full_text), file_path,
    )

    return {
        "text": full_text,
        "num_paragraphs": paragraph_count,
    }


# ===========================================================================
# MAIN ENTRY POINT
# ===========================================================================

def load_document(
    file_path: str,
    min_extracted_chars: int = MIN_EXTRACTED_TEXT_CHARS,
) -> dict:
    """
    Load, validate, extract, and clean text from a resume file.

    This is the primary function callers should invoke. It:
    1. Validates file existence, format (.pdf/.docx), and size limits.
    2. Routes to the appropriate extractor (PDF or DOCX).
    3. Cleans and normalizes extracted text.
    4. Flags suspiciously short extracts without crashing.
    5. Returns a unified dictionary.

    Args:
        file_path: Path to a .pdf or .docx resume file.
        min_extracted_chars: Minimum character threshold to flag suspiciously
                             short output (default: 50 from src.config).

    Returns:
        A dict with:
        - "file_path": Absolute path to the file (as string)
        - "file_type": "pdf" or "docx"
        - "text": Cleaned, normalized extracted text
        - "metadata": Format-specific metadata including:
            - For PDF: {"num_pages": int, "pages": list[str], "is_suspiciously_short": bool}
            - For DOCX: {"num_paragraphs": int, "is_suspiciously_short": bool}

    Raises:
        ValueError: If validation fails (empty path, unsupported type, empty file, too large).
        FileNotFoundError: If the file does not exist.
        RuntimeError: If text extraction fails completely or file is encrypted.
    """
    # Step 1: Validate the file
    path = validate_file(file_path)

    # Step 2: Determine file type and extract
    extension = path.suffix.lower()

    if extension == ".pdf":
        result = extract_text_from_pdf(path)
        extracted_text = result["text"]
        is_suspiciously_short = len(extracted_text.strip()) < min_extracted_chars

        if is_suspiciously_short:
            logger.warning(
                "Extracted text from %s is suspiciously short (%d chars, min expected: %d). "
                "Document may be a scanned image or incomplete.",
                path, len(extracted_text.strip()), min_extracted_chars,
            )

        return {
            "file_path": str(path.resolve()),
            "file_type": "pdf",
            "text": extracted_text,
            "metadata": {
                "num_pages": result["num_pages"],
                "pages": result["pages"],
                "is_suspiciously_short": is_suspiciously_short,
            },
        }

    elif extension == ".docx":
        result = extract_text_from_docx(path)
        extracted_text = result["text"]
        is_suspiciously_short = len(extracted_text.strip()) < min_extracted_chars

        if is_suspiciously_short:
            logger.warning(
                "Extracted text from %s is suspiciously short (%d chars, min expected: %d). "
                "Document may be incomplete.",
                path, len(extracted_text.strip()), min_extracted_chars,
            )

        return {
            "file_path": str(path.resolve()),
            "file_type": "docx",
            "text": extracted_text,
            "metadata": {
                "num_paragraphs": result["num_paragraphs"],
                "is_suspiciously_short": is_suspiciously_short,
            },
        }

    raise ValueError(f"Unsupported file type: {extension}")


# ===========================================================================
# TEXT CHUNKING UTILITY
# ===========================================================================

def chunk_text(
    text: str,
    chunk_size: int = 1000,
    overlap: int = 200,
) -> list[str]:
    """
    Split text into overlapping chunks of a specified size.

    Args:
        text: The text to split into chunks.
        chunk_size: Maximum characters per chunk (default: 1000).
        overlap: Overlapping characters between consecutive chunks (default: 200).

    Returns:
        A list of text chunk strings.

    Raises:
        ValueError: If input text is empty, chunk_size <= 0, overlap < 0,
                    or overlap >= chunk_size.
    """
    if not text or not text.strip():
        raise ValueError("Cannot chunk empty text.")

    if chunk_size <= 0:
        raise ValueError(f"chunk_size must be positive, got {chunk_size}.")

    if overlap < 0:
        raise ValueError(f"overlap must be non-negative, got {overlap}.")

    if overlap >= chunk_size:
        raise ValueError(
            f"overlap ({overlap}) must be less than chunk_size ({chunk_size})."
        )

    if len(text) <= chunk_size:
        return [text]

    chunks = []
    step = chunk_size - overlap
    start = 0

    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end]

        if chunk.strip():
            chunks.append(chunk)

        start += step

    logger.debug(
        "Text chunked: %d characters -> %d chunks (chunk_size=%d, overlap=%d)",
        len(text), len(chunks), chunk_size, overlap,
    )

    return chunks
